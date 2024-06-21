from selenium import webdriver
import openpyxl
wb = openpyxl.load_workbook(filename=r'C:\Users\sipra.das\Downloads\users_data (17).xlsx')
sheet = wb.active
for i in range(1,sheet.max_row+1):
    for j in range(1,sheet.max_column+1):
        print(sheet.cell(i,j).value)

